from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from PIL import Image

from backend.app.infrastructure.conversion.docx_quality import align_docx_page_breaks_to_pdf, apply_docx_heading_styles, apply_docx_page_geometry, apply_docx_section_paper_profiles, apply_docx_toc_navigation, apply_native_list_styles_from_pdf, detect_pdf_page_profiles, enhance_basic_styles, overlay_pdf_visual_layers, promote_repeated_page_header, remove_repeated_page_header_lines, split_docx_inline_outline_titles, split_docx_inline_short_headings, split_docx_multiline_paragraphs, validate_docx_content
from backend.app.infrastructure.conversion.pdf_toc import TocEntry


def test_docx_quality_applies_page_geometry_per_section(tmp_path: Path):
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Pagina 1")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Pagina 2")
    doc.save(str(docx_path))

    apply_docx_page_geometry(docx_path, page_sizes_pt=[(595.0, 842.0), (842.0, 595.0)])

    reloaded = Document(str(docx_path))
    first_section = reloaded.sections[0]
    second_section = reloaded.sections[1]
    assert abs(float(first_section.page_width.pt) - 595.0) < 0.1
    assert abs(float(first_section.page_height.pt) - 842.0) < 0.1
    assert abs(float(second_section.page_width.pt) - 842.0) < 0.1
    assert abs(float(second_section.page_height.pt) - 595.0) < 0.1


def test_docx_quality_detects_standard_paper_profiles():
    profiles = detect_pdf_page_profiles(
        [
            (612.0, 792.0),      # Letter portrait
            (612.0, 1008.0),     # Legal portrait
            (841.89, 595.28),    # A4 landscape
            (700.0, 700.0),      # Custom square
        ]
    )

    assert profiles[0].paper_size == "letter"
    assert profiles[0].orientation == "portrait"
    assert profiles[0].matched_standard is True

    assert profiles[1].paper_size == "legal"
    assert profiles[1].orientation == "portrait"
    assert profiles[1].matched_standard is True

    assert profiles[2].paper_size == "a4"
    assert profiles[2].orientation == "landscape"
    assert profiles[2].matched_standard is True
    assert abs(profiles[2].normalized_width_pt - 841.89) < 0.1
    assert abs(profiles[2].normalized_height_pt - 595.28) < 0.1

    assert profiles[3].paper_size == "custom"
    assert profiles[3].matched_standard is False


def test_docx_quality_applies_section_paper_profiles_when_orientation_is_consistent(tmp_path: Path):
    docx_path = tmp_path / "paper_profiles.docx"
    doc = Document()
    first = doc.sections[0]
    first.page_width = 6000000
    first.page_height = 9000000
    doc.add_paragraph("Pagina 1")
    doc.add_section(WD_SECTION.NEW_PAGE)
    second = doc.sections[1]
    second.page_width = 12000000
    second.page_height = 8000000
    doc.add_paragraph("Pagina 2")
    doc.save(str(docx_path))

    profiles = detect_pdf_page_profiles(
        [
            (612.0, 792.0),
            (792.0, 612.0),
        ]
    )
    apply_docx_section_paper_profiles(docx_path, profiles)


def test_docx_quality_aligns_page_breaks_to_pdf_page_starts(tmp_path: Path):
    import fitz

    docx_path = tmp_path / "anchors.docx"
    doc = Document()
    doc.add_paragraph("Portada")
    doc.add_paragraph("Indice")
    doc.add_paragraph("1. MARCO FORMATIVO")
    doc.add_paragraph("Contenido de la pagina 3")
    doc.save(str(docx_path))

    pdf_path = tmp_path / "anchors.pdf"
    pdf = fitz.open()
    page1 = pdf.new_page()
    page1.insert_text((72, 72), "Portada")
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "Indice")
    page3 = pdf.new_page()
    page3.insert_text((72, 72), "1. MARCO FORMATIVO")
    pdf.save(str(pdf_path))
    pdf.close()

    align_docx_page_breaks_to_pdf(docx_path, pdf_path)

    reloaded = Document(str(docx_path))
    xml = reloaded.paragraphs[1]._p.xml + reloaded.paragraphs[2]._p.xml
    assert 'w:type="page"' in xml

    reloaded = Document(str(docx_path))
    assert abs(float(reloaded.sections[0].page_width.pt) - 612.0) < 0.1
    assert abs(float(reloaded.sections[0].page_height.pt) - 792.0) < 0.1
    assert abs(float(reloaded.sections[1].page_width.pt) - 792.0) < 0.1
    assert abs(float(reloaded.sections[1].page_height.pt) - 612.0) < 0.1


def test_docx_quality_validate_detects_empty_doc(tmp_path: Path):
    docx_path = tmp_path / "empty.docx"
    Document().save(str(docx_path))

    with pytest.raises(ValueError):
        validate_docx_content(docx_path)


def test_docx_quality_converts_literal_bullets_into_word_lists(tmp_path: Path):
    docx_path = tmp_path / "styled.docx"
    doc = Document()
    doc.add_paragraph("●\u200b Primer item")
    doc.add_paragraph("○\u200b Sub item")
    doc.save(str(docx_path))

    enhance_basic_styles(docx_path)

    reloaded = Document(str(docx_path))
    assert reloaded.paragraphs[0].style.name.startswith("List Bullet")
    assert reloaded.paragraphs[0].text == "Primer item"
    assert reloaded.paragraphs[1].style.name.startswith("List Bullet")
    assert reloaded.paragraphs[1].text == "Sub item"


def test_docx_quality_keeps_toc_lines_as_plain_text(tmp_path: Path):
    docx_path = tmp_path / "toc.docx"
    doc = Document()
    doc.add_paragraph("1. MARCO FORMATIVO................................ 3")
    doc.save(str(docx_path))

    enhance_basic_styles(docx_path)

    reloaded = Document(str(docx_path))
    assert reloaded.paragraphs[0].style.name == "Normal"


def test_docx_quality_builds_clickable_toc_links(tmp_path: Path):
    docx_path = tmp_path / "toc_navigation.docx"
    doc = Document()
    doc.add_paragraph("Portada")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Índice")
    doc.add_paragraph("Indice viejo")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Contenido")
    doc.save(str(docx_path))

    apply_docx_toc_navigation(docx_path, [TocEntry(level=1, title="Seccion Uno", page=3)])

    reloaded = Document(str(docx_path))
    index_paragraph = next(p for p in reloaded.paragraphs if "Seccion Uno" in p.text)
    assert "Seccion Uno  3" in index_paragraph.text
    assert 'w:anchor="folio_page_3"' in index_paragraph._p.xml
    content_paragraph = next(p for p in reloaded.paragraphs if p.text == "Contenido")
    assert 'w:name="folio_page_3"' in content_paragraph._p.xml
    assert len(reloaded.sections) == 3


def test_docx_quality_promotes_repeated_page_header(tmp_path: Path):
    docx_path = tmp_path / "header.docx"
    doc = Document()
    doc.add_paragraph("Sistemas Operativos 2")
    doc.add_paragraph("Portada")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Sistemas Operativos 2")
    doc.add_paragraph("Pagina 2")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Sistemas Operativos 2")
    doc.add_paragraph("Pagina 3")
    doc.save(str(docx_path))

    promote_repeated_page_header(docx_path)

    reloaded = Document(str(docx_path))
    body_texts = [paragraph.text for paragraph in reloaded.paragraphs if paragraph.text.strip()]
    assert body_texts.count("Sistemas Operativos 2") == 0
    assert reloaded.sections[0].header.paragraphs[0].text == "Sistemas Operativos 2"


def test_docx_quality_removes_repeated_page_header_without_writing_header(tmp_path: Path):
    docx_path = tmp_path / "header_removed.docx"
    doc = Document()
    doc.add_paragraph("Sistemas Operativos 2")
    doc.add_paragraph("Portada")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Sistemas Operativos 2")
    doc.add_paragraph("Pagina 2")
    doc.save(str(docx_path))

    repeated = remove_repeated_page_header_lines(docx_path)

    reloaded = Document(str(docx_path))
    body_texts = [paragraph.text for paragraph in reloaded.paragraphs if paragraph.text.strip()]
    assert repeated == "Sistemas Operativos 2"
    assert body_texts.count("Sistemas Operativos 2") == 0


def test_docx_quality_applies_heading_styles_from_toc(tmp_path: Path):
    docx_path = tmp_path / "headings.docx"
    doc = Document()
    doc.add_paragraph("2.Enunciado de la Práctica")
    doc.add_paragraph("2.1 Descripción del problema a resolver")
    doc.save(str(docx_path))

    apply_docx_heading_styles(
        docx_path,
        [
            TocEntry(level=1, title="2. Enunciado de la Práctica", page=4),
            TocEntry(level=2, title="2.1 Descripción del problema a resolver", page=4),
        ],
    )

    reloaded = Document(str(docx_path))
    assert reloaded.paragraphs[0].text == "2. Enunciado de la Práctica"
    assert reloaded.paragraphs[0].style.name == "Heading 2"
    assert reloaded.paragraphs[1].style.name == "Heading 3"


def test_docx_quality_splits_multiline_heading_paragraphs(tmp_path: Path):
    docx_path = tmp_path / "multiline.docx"
    doc = Document()
    doc.add_paragraph("2.2.2 Programa intermedio (Daemon en C)\nEste componente será el núcleo del sistema en espacio de usuario.")
    doc.save(str(docx_path))

    split_docx_multiline_paragraphs(
        docx_path,
        [TocEntry(level=2, title="2.2.2 Programa intermedio (Daemon en C)", page=10)],
    )

    reloaded = Document(str(docx_path))
    assert reloaded.paragraphs[0].text == "2.2.2 Programa intermedio (Daemon en C)"
    assert reloaded.paragraphs[1].text == "Este componente será el núcleo del sistema en espacio de usuario."


def test_docx_quality_splits_inline_outline_titles(tmp_path: Path):
    docx_path = tmp_path / "inline_outline.docx"
    doc = Document()
    doc.add_paragraph("Visualización de métricas (obligatorio) El dashboard deberá reutilizar las gráficas definidas en la Práctica 6.")
    doc.save(str(docx_path))

    split_docx_inline_outline_titles(
        docx_path,
        [TocEntry(level=3, title="Visualización de métricas (obligatorio)", page=14)],
    )

    reloaded = Document(str(docx_path))
    assert reloaded.paragraphs[0].text == "Visualización de métricas (obligatorio)"
    assert reloaded.paragraphs[1].text.startswith("El dashboard deberá reutilizar")


def test_docx_quality_splits_inline_short_headings(tmp_path: Path):
    docx_path = tmp_path / "short_heading.docx"
    doc = Document()
    doc.add_paragraph("Funciones principales El programa deberá invocar periódicamente las syscalls.")
    doc.save(str(docx_path))

    split_docx_inline_short_headings(docx_path)

    reloaded = Document(str(docx_path))
    assert reloaded.paragraphs[0].text == "Funciones principales"
    assert reloaded.paragraphs[0].style.name == "Heading 4"
    assert reloaded.paragraphs[1].text.startswith("El programa deberá invocar")


def test_docx_quality_overlays_pdf_visual_layers(tmp_path: Path):
    import fitz

    pdf_path = tmp_path / "sample.pdf"
    docx_path = tmp_path / "sample.docx"

    image = Image.new("RGB", (80, 40), color=(255, 140, 0))
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")

    pdf = fitz.open()
    page = pdf.new_page(width=300, height=400)
    page.insert_image(fitz.Rect(50, 90, 130, 130), stream=image_bytes.getvalue())
    shape = page.new_shape()
    shape.draw_rect(fitz.Rect(30, 60, 180, 80))
    shape.finish(color=(0, 0, 0), fill=(0.8, 0.8, 0.8), width=1)
    shape.commit()
    pdf.save(str(pdf_path))
    pdf.close()

    doc = Document()
    doc.add_paragraph("Portada")
    doc.save(str(docx_path))

    overlay_pdf_visual_layers(docx_path, pdf_path)

    reloaded = Document(str(docx_path))
    assert "wp:anchor" in reloaded.paragraphs[0]._p.xml


def test_docx_quality_applies_native_list_styles_from_pdf(tmp_path: Path):
    import fitz

    pdf_path = tmp_path / "lists.pdf"
    docx_path = tmp_path / "lists.docx"

    pdf = fitz.open()
    page = pdf.new_page(width=300, height=400)
    page.insert_text((40, 60), "● Primer item", fontsize=11)
    page.insert_text((60, 85), "○ Sub item", fontsize=11)
    pdf.save(str(pdf_path))
    pdf.close()

    doc = Document()
    doc.add_paragraph("● Primer item")
    doc.add_paragraph("○ Sub item")
    doc.save(str(docx_path))

    apply_native_list_styles_from_pdf(docx_path, pdf_path)

    reloaded = Document(str(docx_path))
    assert reloaded.paragraphs[0].text == "Primer item"
    assert reloaded.paragraphs[0].style.name.startswith("List Bullet")
    assert reloaded.paragraphs[1].text == "Sub item"
    assert reloaded.paragraphs[1].style.name.startswith("List Bullet")
