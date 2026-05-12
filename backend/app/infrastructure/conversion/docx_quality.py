from __future__ import annotations

import copy
from dataclasses import dataclass
from io import BytesIO
import re
from collections import Counter
from pathlib import Path
from typing import TYPE_CHECKING

from PIL import Image, ImageDraw

from backend.app.infrastructure.conversion.pdf_toc import clean_toc_title, extract_pdf_toc_entries, normalize_lookup
from backend.app.infrastructure.conversion.text_cleaner import repair_common_extraction_typos

if TYPE_CHECKING:
    from backend.app.infrastructure.conversion.pdf_toc import TocEntry


_BULLET_MARKERS = "•●○▪■◦"
_CONTROL_CHARS_RE = re.compile(r"[\u200b\u200c\u200d\ufeff]")
_BULLET_RE = re.compile(rf"^\s*([{re.escape(_BULLET_MARKERS)}])\s*")
_NUMBER_RE = re.compile(r"^\s*((?:\d+|[A-Za-z]))([.)])\s+")
_INLINE_LIST_SPLIT_RE = re.compile(rf"\s+(?=(?:[{re.escape(_BULLET_MARKERS)}]\s*|\d+[.)]\s+|[A-Za-z][.)]\s+))")
_STANDARD_PAPER_SIZES_PT: dict[str, tuple[float, float]] = {
    "letter": (612.0, 792.0),
    "legal": (612.0, 1008.0),
    "a4": (595.28, 841.89),
    "tabloid": (792.0, 1224.0),
    "a3": (841.89, 1190.55),
}

_ANCHOR_TOKEN_RE = re.compile(r"[^0-9a-záéíóúüñ]+", re.IGNORECASE)


@dataclass(frozen=True)
class PdfPageProfile:
    width_pt: float
    height_pt: float
    normalized_width_pt: float
    normalized_height_pt: float
    paper_size: str
    orientation: str
    matched_standard: bool


def _clean_text_fragment(text: str) -> str:
    text = _CONTROL_CHARS_RE.sub("", text or "")
    text = text.replace("\t", " ")
    text = re.sub(r" {2,}", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = repair_common_extraction_typos(text)
    return text.strip()


def _normalize_anchor_text(text: str) -> str:
    cleaned = _clean_text_fragment(text).lower()
    cleaned = _ANCHOR_TOKEN_RE.sub(" ", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned.strip()


def _rewrite_paragraph_text(paragraph, text: str) -> None:
    cleaned = _clean_text_fragment(text)
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = cleaned
    elif cleaned:
        paragraph.add_run(cleaned)


def _insert_paragraph_after(paragraph):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    new_paragraph = OxmlElement("w:p")
    paragraph._p.addnext(new_paragraph)
    return Paragraph(new_paragraph, paragraph._parent)


def _paragraph_segments(text: str) -> list[str]:
    normalized = _clean_text_fragment(text)
    if not normalized:
        return []

    segments: list[str] = []
    for raw_line in normalized.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        parts = [part.strip() for part in _INLINE_LIST_SPLIT_RE.split(line) if part.strip()]
        segments.extend(parts or [line])
    return segments


def _meaningful_lines(text: str) -> list[str]:
    return [_clean_text_fragment(line) for line in text.splitlines() if _clean_text_fragment(line)]


def _extract_toc_lines(text: str) -> list[str]:
    normalized = _clean_text_fragment(text).replace("\t", " ")
    normalized = re.sub(r"\s{2,}", " ", normalized)
    matches = [match.group(0).strip() for match in re.finditer(r".+?\.{3,}\s*\d+", normalized)]
    return matches


def _parse_list_segment(segment: str) -> tuple[str, str | None, str]:
    if "..." in segment or "……" in segment:
        return "text", None, _clean_text_fragment(segment)

    bullet_match = _BULLET_RE.match(segment)
    if bullet_match:
        marker = bullet_match.group(1)
        return "bullet", marker, _clean_text_fragment(segment[bullet_match.end():])

    number_match = _NUMBER_RE.match(segment)
    if number_match:
        return "number", number_match.group(1), _clean_text_fragment(segment[number_match.end():])

    return "text", None, _clean_text_fragment(segment)


def _looks_like_toc_entry(text: str) -> bool:
    cleaned = _clean_text_fragment(text)
    return bool(re.search(r"\.{5,}\s*\d+\s*$", cleaned))


def _list_style_name(document, kind: str, level: int) -> str | None:
    level = max(1, min(level, 3))
    if kind == "bullet":
        candidates = [f"List Bullet {level}", "List Bullet"] if level > 1 else ["List Bullet"]
    else:
        candidates = [f"List Number {level}", "List Number"] if level > 1 else ["List Number"]
    return _pick_style(document, candidates)


def _apply_native_list_style(paragraph, document, kind: str, level: int, text: str, *, rewrite_text: bool) -> None:
    if rewrite_text:
        _rewrite_paragraph_text(paragraph, text)
    style_name = _list_style_name(document, kind, level)
    if style_name:
        paragraph.style = document.styles[style_name]
    try:
        from docx.shared import Inches, Pt

        base_indent = 0.30 + (max(level, 1) - 1) * 0.24
        paragraph.paragraph_format.left_indent = Inches(base_indent)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
    except Exception:
        pass


def _extract_pdf_native_list_candidates(pdf_path: Path) -> dict[str, tuple[str, int]]:
    import fitz

    candidates: dict[str, tuple[str, int]] = {}
    pdf = fitz.open(str(pdf_path))
    try:
        for page in pdf:
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    spans = [span for span in line.get("spans", []) if span.get("text")]
                    if not spans:
                        continue
                    raw_text = "".join(span.get("text", "") for span in spans)
                    cleaned = _clean_text_fragment(raw_text)
                    if not cleaned or _looks_like_toc_entry(cleaned):
                        continue

                    kind, marker, text = _parse_list_segment(cleaned)
                    if kind == "text" or not text:
                        continue

                    max_size = max(float(span.get("size") or 0.0) for span in spans)
                    # Avoid turning section titles into list items.
                    if kind == "number" and max_size >= 14:
                        continue

                    level = 1
                    if kind == "bullet":
                        if marker in {"○", "◦"}:
                            level = 2
                        elif marker in {"▪", "■"}:
                            level = 3
                    elif marker and marker.isalpha():
                        level = 2

                    normalized = normalize_lookup(text)
                    if normalized and normalized not in candidates:
                        candidates[normalized] = (kind, level)
    finally:
        pdf.close()

    return candidates


def _pick_style(document, style_names: list[str]) -> str | None:
    for style_name in style_names:
        try:
            document.styles[style_name]
            return style_name
        except KeyError:
            continue
    return None


def _apply_block(paragraph, document, kind: str, marker: str | None, text: str) -> None:
    _rewrite_paragraph_text(paragraph, text)

    if kind == "bullet":
        if marker in {"○", "◦"}:
            style_name = _pick_style(document, ["List Bullet 2", "List Bullet"])
        elif marker in {"▪", "■"}:
            style_name = _pick_style(document, ["List Bullet 3", "List Bullet 2", "List Bullet"])
        else:
            style_name = _pick_style(document, ["List Bullet"])
        if style_name:
            paragraph.style = document.styles[style_name]
        return

    if kind == "number":
        style_name = _pick_style(document, ["List Number"])
        if style_name:
            paragraph.style = document.styles[style_name]
        return

    style_name = _pick_style(document, ["Normal"])
    if style_name and ((paragraph.style.name if paragraph.style else "") or "").startswith("List "):
        paragraph.style = document.styles[style_name]


def _delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _clear_paragraph_content(paragraph) -> None:
    paragraph_element = paragraph._p
    for child in list(paragraph_element):
        if child.tag.endswith("}pPr"):
            continue
        paragraph_element.remove(child)


def _copy_section_properties(from_paragraph, to_paragraph) -> None:
    paragraph_properties = from_paragraph._p.get_or_add_pPr()
    section_properties = paragraph_properties.sectPr
    if section_properties is None:
        return
    target_properties = to_paragraph._p.get_or_add_pPr()
    existing = target_properties.sectPr
    if existing is not None:
        target_properties.remove(existing)
    target_properties.append(copy.deepcopy(section_properties))


def _page_start_indices(paragraphs: list) -> list[int]:
    if not paragraphs:
        return []

    starts = [0]
    for index, paragraph in enumerate(paragraphs[:-1]):
        paragraph_properties = paragraph._p.pPr
        if paragraph_properties is not None and paragraph_properties.sectPr is not None:
            starts.append(index + 1)
    return starts


def _append_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), str(bookmark_id))
    bookmark_start.set(qn("w:name"), name)

    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), str(bookmark_id))

    paragraph._p.insert(0, bookmark_start)
    paragraph._p.append(bookmark_end)


def _append_internal_hyperlink(paragraph, text: str, anchor: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_style = OxmlElement("w:rStyle")
    run_style.set(qn("w:val"), "Hyperlink")
    run_properties.append(run_style)
    run.append(run_properties)

    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)

    paragraph._p.append(hyperlink)
    return hyperlink


def _append_internal_hyperlink_plain(paragraph, text: str, anchor: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)

    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _append_floating_picture(paragraph, image_bytes: bytes, x_pt: float, y_pt: float, width_pt: float, height_pt: float, shape_id: int, behind_text: bool = True) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run = paragraph.add_run()
    run.add_picture(BytesIO(image_bytes), width=Pt(width_pt), height=Pt(height_pt))

    drawing = run._r.xpath("./w:drawing")
    if not drawing:
        return
    drawing = drawing[0]
    inline_nodes = drawing.xpath("./wp:inline")
    if not inline_nodes:
        return
    inline = inline_nodes[0]

    extent = inline.xpath("./wp:extent")[0]
    graphic = inline.xpath("./a:graphic")[0]
    c_nv = inline.xpath("./wp:cNvGraphicFramePr")

    anchor = OxmlElement("wp:anchor")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "0")
    anchor.set("distR", "0")
    anchor.set("simplePos", "0")
    anchor.set("relativeHeight", "251659264")
    anchor.set("behindDoc", "1" if behind_text else "0")
    anchor.set("locked", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "1")

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "page")
    pos_offset_h = OxmlElement("wp:posOffset")
    pos_offset_h.text = str(int(round(x_pt * 12700)))
    position_h.append(pos_offset_h)
    anchor.append(position_h)

    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "page")
    pos_offset_v = OxmlElement("wp:posOffset")
    pos_offset_v.text = str(int(round(y_pt * 12700)))
    position_v.append(pos_offset_v)
    anchor.append(position_v)

    anchor.append(copy.deepcopy(extent))

    effect_extent = OxmlElement("wp:effectExtent")
    effect_extent.set("l", "0")
    effect_extent.set("t", "0")
    effect_extent.set("r", "0")
    effect_extent.set("b", "0")
    anchor.append(effect_extent)

    wrap_none = OxmlElement("wp:wrapNone")
    anchor.append(wrap_none)

    doc_pr = OxmlElement("wp:docPr")
    doc_pr.set("id", str(shape_id))
    doc_pr.set("name", f"FolioExtractImage{shape_id}")
    anchor.append(doc_pr)

    if c_nv:
        anchor.append(copy.deepcopy(c_nv[0]))
    else:
        anchor.append(OxmlElement("wp:cNvGraphicFramePr"))

    anchor.append(copy.deepcopy(graphic))

    drawing.remove(inline)
    drawing.append(anchor)


def _font_family_name(raw_name: str | None) -> str:
    if not raw_name:
        return "Arial"
    name = raw_name.replace("MT", "").strip()
    name = name.split("-")[0].strip()
    return name or "Arial"


def _color_int_to_hex(color: int | None) -> str | None:
    if color is None:
        return None
    return f"{int(color) & 0xFFFFFF:06X}"


def _shape_style(x_pt: float, y_pt: float, width_pt: float, height_pt: float, z_index: int) -> str:
    return (
        "position:absolute;"
        f"margin-left:{x_pt:.2f}pt;"
        f"margin-top:{y_pt:.2f}pt;"
        f"width:{max(width_pt, 1.0):.2f}pt;"
        f"height:{max(height_pt, 1.0):.2f}pt;"
        f"z-index:{z_index};"
        "mso-position-horizontal:absolute;"
        "mso-position-vertical:absolute;"
        "mso-position-horizontal-relative:page;"
        "mso-position-vertical-relative:page;"
    )


def _append_positioned_textbox(paragraph, x_pt: float, y_pt: float, width_pt: float, height_pt: float, line: dict, shape_id: int) -> None:
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn

    spans = [span for span in line.get("spans", []) if span.get("text")]
    if not spans:
        return

    pict_run = paragraph.add_run()
    pict = OxmlElement("w:pict")

    shape = parse_xml(
        '<v:shape xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:o="urn:schemas-microsoft-com:office:office"/>'
    )
    shape.set("{urn:schemas-microsoft-com:office:office}spid", f"_x0000_s{shape_id}")
    shape.set("id", f"FolioExtractTextBox{shape_id}")
    shape.set("stroked", "f")
    shape.set("filled", "f")
    shape.set("style", _shape_style(x_pt, y_pt, width_pt, height_pt, 300000 + shape_id))

    textbox = parse_xml('<v:textbox xmlns:v="urn:schemas-microsoft-com:vml"/>')
    textbox.set("inset", "0,0,0,0")
    textbox.set("style", "mso-fit-shape-to-text:t")
    txbx_content = OxmlElement("w:txbxContent")

    text_paragraph = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    p_pr.append(spacing)
    text_paragraph.append(p_pr)

    for span in spans:
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")

        font_name = _font_family_name(span.get("font"))
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), font_name)
        r_fonts.set(qn("w:hAnsi"), font_name)
        r_fonts.set(qn("w:cs"), font_name)
        r_pr.append(r_fonts)

        font_size = max(1, int(round(float(span.get("size") or 11.0) * 2)))
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(font_size))
        r_pr.append(sz)
        sz_cs = OxmlElement("w:szCs")
        sz_cs.set(qn("w:val"), str(font_size))
        r_pr.append(sz_cs)

        raw_font = span.get("font") or ""
        flags = int(span.get("flags") or 0)
        if "Bold" in raw_font or (flags & 16):
            r_pr.append(OxmlElement("w:b"))
        if "Italic" in raw_font or (flags & 2):
            r_pr.append(OxmlElement("w:i"))

        color_hex = _color_int_to_hex(span.get("color"))
        if color_hex:
            color = OxmlElement("w:color")
            color.set(qn("w:val"), color_hex)
            r_pr.append(color)

        run.append(r_pr)
        text_node = OxmlElement("w:t")
        if span.get("text", "").startswith(" ") or span.get("text", "").endswith(" "):
            text_node.set(qn("xml:space"), "preserve")
        text_node.text = span.get("text", "")
        run.append(text_node)
        text_paragraph.append(run)

    txbx_content.append(text_paragraph)
    textbox.append(txbx_content)
    shape.append(textbox)
    pict.append(shape)
    pict_run._r.append(pict)


def _page_anchor_paragraphs(document) -> list:
    paragraphs = list(document.paragraphs)
    if not paragraphs:
        return []
    return [paragraphs[index] for index in _page_start_indices(paragraphs) if index < len(paragraphs)]


def _insert_page_break_before(paragraph) -> None:
    from docx.enum.text import WD_BREAK

    before = paragraph.insert_paragraph_before()
    before.add_run().add_break(WD_BREAK.PAGE)


def _insert_page_break_before_table(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    page_break = OxmlElement("w:br")
    page_break.set(qn("w:type"), "page")
    run.append(page_break)
    paragraph.append(run)
    table._element.addprevious(paragraph)


def _iter_document_blocks(document):
    from docx.document import Document as DocxDocument
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent_elm = document.element.body if isinstance(document, DocxDocument) else document._element
    parent = document if isinstance(document, DocxDocument) else document._parent

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _block_text(block) -> str:
    paragraph_text = getattr(block, "text", "")
    if paragraph_text:
        return paragraph_text

    rows = getattr(block, "rows", None)
    if rows is None:
        return ""

    parts: list[str] = []
    for row in rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                text = _clean_text_fragment(paragraph.text)
                if text:
                    parts.append(text)
            if len(parts) >= 6:
                return " ".join(parts)
    return " ".join(parts)


def _pdf_page_anchor_lines(pdf_path: Path, *, max_candidates_per_page: int = 6) -> list[list[str]]:
    import fitz

    pdf = fitz.open(str(pdf_path))
    try:
        page_lines: list[list[str]] = []
        repeated_counter: Counter[str] = Counter()

        raw_pages: list[list[str]] = []
        for page in pdf:
            lines = [
                _clean_text_fragment(line)
                for line in page.get_text("text").splitlines()
                if _clean_text_fragment(line)
            ]
            raw_pages.append(lines)
            for line in lines[:3]:
                normalized = _normalize_anchor_text(line)
                if normalized:
                    repeated_counter[normalized] += 1

        for lines in raw_pages:
            candidates: list[str] = []
            for line in lines:
                normalized = _normalize_anchor_text(line)
                if not normalized:
                    continue
                if repeated_counter.get(normalized, 0) > max(2, len(raw_pages) // 3):
                    continue
                if len(normalized) < 4:
                    continue
                if normalized in candidates:
                    continue
                candidates.append(normalized)
                if len(candidates) >= max_candidates_per_page:
                    break
            page_lines.append(candidates)

        return page_lines
    finally:
        pdf.close()


def align_docx_page_breaks_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    """
    Force DOCX page starts to follow the source PDF's page starts as closely as possible
    by inserting page breaks before paragraphs matching the first meaningful lines of each PDF page.
    """
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(docx_path))
    blocks = list(_iter_document_blocks(document))
    if not blocks:
        return

    anchor_candidates = _pdf_page_anchor_lines(pdf_path)
    if len(anchor_candidates) <= 1:
        return

    block_norms = [_normalize_anchor_text(_block_text(block)) for block in blocks]
    last_match_index = 0
    changed = False

    for page_number in range(1, len(anchor_candidates)):
        candidates = anchor_candidates[page_number]
        if not candidates:
            continue

        match_index: int | None = None
        for candidate in candidates:
            for idx in range(last_match_index + 1, len(blocks)):
                block_text = block_norms[idx]
                if not block_text:
                    continue
                if candidate == block_text or candidate in block_text or block_text in candidate:
                    match_index = idx
                    break
            if match_index is not None:
                break

        if match_index is None or match_index <= 0:
            continue

        current_block = blocks[match_index]
        current_xml = current_block._element.xml
        previous_xml = blocks[match_index - 1]._element.xml
        if 'w:type="page"' in previous_xml or 'w:type="page"' in current_xml:
            last_match_index = match_index
            continue

        if isinstance(current_block, Paragraph):
            _insert_page_break_before(current_block)
        elif isinstance(current_block, Table):
            _insert_page_break_before_table(current_block)
        else:
            continue
        changed = True
        blocks = list(_iter_document_blocks(document))
        block_norms = [_normalize_anchor_text(_block_text(block)) for block in blocks]
        last_match_index = match_index + 1

    if changed:
        document.save(str(docx_path))


def _toc_title_match_keys(title: str) -> list[str]:
    keys: list[str] = []
    normalized = normalize_lookup(title)
    if normalized:
        keys.append(normalized)

    stripped = re.sub(r"^(?:\d+(?:\.\d+)*|[A-Za-z])\s+", "", clean_toc_title(title))
    stripped_normalized = normalize_lookup(stripped)
    if stripped_normalized and stripped_normalized not in keys:
        keys.append(stripped_normalized)

    return keys


def _preview_toc_page_numbers(preview_pdf_path: Path, toc_entries: list["TocEntry"]) -> dict[str, int]:
    import fitz

    preview = fitz.open(str(preview_pdf_path))
    try:
        page_texts = [normalize_lookup(page.get_text("text")) for page in preview]
    finally:
        preview.close()

    page_numbers: dict[str, int] = {}
    for entry in toc_entries:
        keys = _toc_title_match_keys(entry.title)
        if not keys:
            continue
        for page_number, page_text in enumerate(page_texts, start=1):
            if any(key and key in page_text for key in keys):
                page_numbers[normalize_lookup(entry.title)] = page_number
                break
    return page_numbers


def _insert_page_break_before_heading_match(docx_path: Path, title: str) -> bool:
    from docx import Document

    title_keys = _toc_title_match_keys(title)
    if not title_keys:
        return False

    document = Document(str(docx_path))
    target_paragraph = None

    for paragraph in document.paragraphs:
        text = _clean_text_fragment(paragraph.text)
        if not text or _looks_like_toc_entry(text):
            continue
        normalized = normalize_lookup(text)
        if not normalized or normalized == "indice":
            continue
        if any(key == normalized or key in normalized or normalized in key for key in title_keys):
            target_paragraph = paragraph
            break

    if target_paragraph is None:
        return False

    previous_xml = target_paragraph._p.getprevious()
    if previous_xml is not None and 'w:type="page"' in previous_xml.xml:
        return False

    if 'w:type="page"' in target_paragraph._p.xml:
        return False

    _insert_page_break_before(target_paragraph)
    document.save(str(docx_path))
    return True


def align_docx_toc_pages_with_preview(docx_path: Path, source_pdf_path: Path, preview_pdf_path: Path) -> int:
    """
    Compare PDF TOC page numbers against the DOCX preview PDF and insert a small number
    of page breaks before major headings that have drifted to an earlier page.
    """
    toc_entries = extract_pdf_toc_entries(source_pdf_path)
    if not toc_entries:
        return 0

    top_level = min(entry.level for entry in toc_entries)
    major_entries = [entry for entry in toc_entries if entry.level == top_level and entry.page > 1]
    if not major_entries:
        return 0

    preview_page_numbers = _preview_toc_page_numbers(preview_pdf_path, major_entries)
    lagging_entries = [
        entry
        for entry in major_entries
        if preview_page_numbers.get(normalize_lookup(entry.title), entry.page) < entry.page
    ]
    if not lagging_entries:
        return 0

    inserted = 0
    for entry in sorted(lagging_entries, key=lambda item: item.page, reverse=True):
        if _insert_page_break_before_heading_match(docx_path, entry.title):
            inserted += 1

    return inserted


def _iter_pdf_image_blocks(pdf_path: Path):
    import fitz

    pdf = fitz.open(str(pdf_path))
    try:
        for page_index, page in enumerate(pdf):
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 1:
                    continue
                x0, y0, x1, y1 = block["bbox"]
                width = x1 - x0
                height = y1 - y0
                image_bytes = block.get("image")
                if not image_bytes:
                    continue
                yield page_index, image_bytes, float(x0), float(y0), float(width), float(height)
    finally:
        pdf.close()


def _rgb255(rgb: tuple[float, float, float] | None, alpha: float | None = 1.0) -> tuple[int, int, int, int]:
    if rgb is None:
        return (0, 0, 0, 0)
    a = int(max(0.0, min(1.0, alpha if alpha is not None else 1.0)) * 255)
    return tuple(int(max(0.0, min(1.0, channel)) * 255) for channel in rgb) + (a,)


def _render_page_visual_layer(page, scale: float = 2.0) -> bytes | None:
    width = max(1, int(round(float(page.rect.width) * scale)))
    height = max(1, int(round(float(page.rect.height) * scale)))
    canvas = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    draw = ImageDraw.Draw(canvas)
    has_visual_content = False

    for drawing in page.get_drawings():
        rect = drawing.get("rect")
        fill = drawing.get("fill")
        color = drawing.get("color")
        draw_type = drawing.get("type")

        # Skip the full-page white background rectangle.
        if (
            draw_type == "f"
            and fill == (1.0, 1.0, 1.0)
            and rect is not None
            and rect.x0 <= 1
            and rect.y0 <= 1
            and abs(rect.x1 - page.rect.width) <= 1.5
            and abs(rect.y1 - page.rect.height) <= 1.5
        ):
            continue

        fill_color = _rgb255(fill, drawing.get("fill_opacity"))
        stroke_color = _rgb255(color, drawing.get("stroke_opacity"))
        line_width = max(1, int(round(float(drawing.get("width") or 1.0) * scale)))

        for item in drawing.get("items", []):
            operator = item[0]
            if operator == "re":
                rect = item[1]
                box = [rect.x0 * scale, rect.y0 * scale, rect.x1 * scale, rect.y1 * scale]
                if draw_type in {"f", "fs"} and fill is not None:
                    draw.rectangle(box, fill=fill_color)
                    has_visual_content = True
                if draw_type in {"s", "fs"} and color is not None:
                    draw.rectangle(box, outline=stroke_color, width=line_width)
                    has_visual_content = True
            elif operator == "l":
                p1, p2 = item[1], item[2]
                draw.line((p1.x * scale, p1.y * scale, p2.x * scale, p2.y * scale), fill=stroke_color, width=line_width)
                has_visual_content = True

    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 1 or not block.get("image"):
            continue
        x0, y0, x1, y1 = block["bbox"]
        try:
            block_image = Image.open(BytesIO(block["image"])).convert("RGBA")
        except Exception:
            continue
        dest_width = max(1, int(round((x1 - x0) * scale)))
        dest_height = max(1, int(round((y1 - y0) * scale)))
        if block_image.size != (dest_width, dest_height):
            block_image = block_image.resize((dest_width, dest_height), Image.Resampling.LANCZOS)
        canvas.alpha_composite(block_image, (int(round(x0 * scale)), int(round(y0 * scale))))
        has_visual_content = True

    if not has_visual_content:
        return None

    output = BytesIO()
    canvas.save(output, format="PNG")
    return output.getvalue()


def _page_line_items(page) -> list[dict]:
    items: list[dict] = []
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = [span for span in line.get("spans", []) if span.get("text")]
            if not spans:
                continue
            line_text = "".join(span.get("text", "") for span in spans)
            if not line_text.strip():
                continue
            items.append(
                {
                    "bbox": tuple(float(v) for v in line["bbox"]),
                    "spans": spans,
                }
            )
    return items


def split_docx_multiline_paragraphs(docx_path: Path, toc_entries: list["TocEntry"]) -> None:
    """Split extracted paragraphs that contain multiple semantic blocks separated by newlines."""
    from docx import Document

    document = Document(str(docx_path))
    toc_titles = {normalize_lookup(entry.title) for entry in toc_entries}

    for paragraph in list(document.paragraphs):
        raw_text = paragraph.text or ""
        if "\n" not in raw_text:
            continue
        if "w:hyperlink" in paragraph._p.xml:
            continue

        lines = _meaningful_lines(raw_text)
        if len(lines) < 2:
            continue

        first_line = lines[0]
        should_split = normalize_lookup(first_line) in toc_titles
        if not should_split:
            should_split = any(
                re.match(r"^(?:\d+[.)]?|Thread \d|[A-Z]\.|[•●○▪■◦-])", line)
                or line.endswith(":")
                for line in lines[1:]
            )
        if not should_split:
            continue

        _rewrite_paragraph_text(paragraph, first_line)
        cursor = paragraph
        for line in lines[1:]:
            cursor = _insert_paragraph_after(cursor)
            _rewrite_paragraph_text(cursor, line)

    document.save(str(docx_path))


def split_docx_inline_outline_titles(docx_path: Path, toc_entries: list["TocEntry"]) -> None:
    """Split paragraphs that start with a known outline title followed by body text."""
    from docx import Document

    document = Document(str(docx_path))
    ordered_titles = sorted(
        {clean_toc_title(entry.title) for entry in toc_entries},
        key=len,
        reverse=True,
    )

    for paragraph in list(document.paragraphs):
        text = _clean_text_fragment(paragraph.text)
        if not text or "w:hyperlink" in paragraph._p.xml or "\t" in text:
            continue

        for title in ordered_titles:
            if not text.startswith(title):
                continue
            remainder = _clean_text_fragment(text[len(title):])
            if not remainder or len(remainder) < 20:
                continue
            if remainder.startswith((".", ":", "-", ",")):
                continue
            _rewrite_paragraph_text(paragraph, title)
            follower = _insert_paragraph_after(paragraph)
            _rewrite_paragraph_text(follower, remainder)
            break

    document.save(str(docx_path))


def split_docx_inline_short_headings(docx_path: Path) -> None:
    """Split short inline heading phrases from the explanatory sentence that follows."""
    from docx import Document

    document = Document(str(docx_path))
    pattern = re.compile(
        r"^(?P<title>(?:[A-ZÁÉÍÓÚÜÑ][\wÁÉÍÓÚÜÑáéíóúüñ/()+.-]*\s+){0,5}[A-ZÁÉÍÓÚÜÑ][\wÁÉÍÓÚÜÑáéíóúüñ/()+.-]*)(?P<rest>\s+(?:El|La|Los|Las|Para|Se|Puede|Debe|Esta|Estas|Este|Estos)\b.*)$"
    )

    for paragraph in list(document.paragraphs):
        text = _clean_text_fragment(paragraph.text)
        if not text or "w:hyperlink" in paragraph._p.xml or "\t" in text:
            continue
        if (paragraph.style.name if paragraph.style else "").startswith("Heading"):
            continue

        match = pattern.match(text)
        if not match:
            continue

        title = _clean_text_fragment(match.group("title"))
        rest = _clean_text_fragment(match.group("rest"))
        if len(title.split()) > 6 or len(title) > 70 or len(rest) < 20:
            continue

        _rewrite_paragraph_text(paragraph, title)
        try:
            paragraph.style = document.styles["Heading 4"]
        except KeyError:
            pass
        follower = _insert_paragraph_after(paragraph)
        _rewrite_paragraph_text(follower, rest)

    document.save(str(docx_path))


def promote_repeated_page_header(docx_path: Path) -> None:
    """Move a repeated top-of-page line into the actual Word header and remove body duplicates."""
    from docx import Document

    document = Document(str(docx_path))
    paragraphs = list(document.paragraphs)
    page_starts = _page_start_indices(paragraphs)
    if not page_starts:
        return

    first_lines = [
        _clean_text_fragment(paragraphs[index].text)
        for index in page_starts
        if index < len(paragraphs) and _clean_text_fragment(paragraphs[index].text)
    ]
    if not first_lines:
        return

    repeated_line, frequency = Counter(
        line for line in first_lines if 2 <= len(line) <= 120
    ).most_common(1)[0]
    if frequency < max(3, len(page_starts) // 2):
        return

    for section in document.sections:
        header = section.header
        header.is_linked_to_previous = False
        header_paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        _rewrite_paragraph_text(header_paragraph, repeated_line)
        for extra_paragraph in list(header.paragraphs[1:]):
            _delete_paragraph(extra_paragraph)

    for page_start in reversed(page_starts):
        if page_start < len(paragraphs) and _clean_text_fragment(paragraphs[page_start].text) == repeated_line:
            _delete_paragraph(paragraphs[page_start])

    document.save(str(docx_path))


def remove_repeated_page_header_lines(docx_path: Path) -> str | None:
    """Remove a repeated top-of-page line from body paragraphs without writing any visible header text."""
    from docx import Document

    document = Document(str(docx_path))
    paragraphs = list(document.paragraphs)
    page_starts = _page_start_indices(paragraphs)
    if not page_starts:
        return None

    first_lines = [
        _clean_text_fragment(paragraphs[index].text)
        for index in page_starts
        if index < len(paragraphs) and _clean_text_fragment(paragraphs[index].text)
    ]
    if not first_lines:
        return None

    repeated_line, frequency = Counter(
        line for line in first_lines if 2 <= len(line) <= 120
    ).most_common(1)[0]
    if frequency < max(3, len(page_starts) // 2):
        return None

    for page_start in reversed(page_starts):
        if page_start < len(paragraphs) and _clean_text_fragment(paragraphs[page_start].text) == repeated_line:
            _delete_paragraph(paragraphs[page_start])

    document.save(str(docx_path))
    return repeated_line


def apply_docx_toc_navigation(docx_path: Path, toc_entries: list["TocEntry"]) -> None:
    """Replace a dead index with clickable internal links to the matching page bookmarks."""
    if not toc_entries:
        return

    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Inches, Pt

    document = Document(str(docx_path))
    paragraphs = list(document.paragraphs)
    if not paragraphs:
        return

    page_starts = _page_start_indices(paragraphs)
    page_bookmarks: dict[int, str] = {}
    bookmark_id = 1
    for page_number, paragraph_index in enumerate(page_starts, start=1):
        bookmark_name = f"folio_page_{page_number}"
        _append_bookmark(paragraphs[paragraph_index], bookmark_name, bookmark_id)
        page_bookmarks[page_number] = bookmark_name
        bookmark_id += 1

    toc_index = next(
        (
            idx
            for idx, paragraph in enumerate(paragraphs)
            if _clean_text_fragment(paragraph.text).lower() in {"indice", "índice"}
        ),
        None,
    )
    if toc_index is None:
        document.save(str(docx_path))
        return

    toc_page_index = max(i for i, start in enumerate(page_starts) if start <= toc_index)
    toc_end = page_starts[toc_page_index + 1] if toc_page_index + 1 < len(page_starts) else len(paragraphs)
    section_break_paragraph = paragraphs[toc_end - 1] if toc_end - 1 > toc_index else None

    for index in range(toc_end - 1, toc_index, -1):
        _delete_paragraph(paragraphs[index])

    toc_paragraph = paragraphs[toc_index]
    cursor = toc_paragraph
    for entry in toc_entries:
        target = page_bookmarks.get(entry.page)
        if not target:
            continue
        cursor = _insert_paragraph_after(cursor)
        cursor.paragraph_format.left_indent = Inches(0.2 * max(0, entry.level - 1))
        cursor.paragraph_format.first_line_indent = Pt(0)
        cursor.paragraph_format.space_after = Pt(0)
        cursor.paragraph_format.space_before = Pt(0)
        tab_stop = Inches(6.15)
        cursor.paragraph_format.tab_stops.add_tab_stop(tab_stop, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        title = clean_toc_title(entry.title)
        _append_internal_hyperlink(cursor, title, target)
        if entry.level == 1 and cursor.runs:
            cursor.runs[-1].bold = True
        cursor.add_run(f"\t{entry.page}")

    if section_break_paragraph is not None:
        _copy_section_properties(section_break_paragraph, cursor)

    document.save(str(docx_path))


def apply_docx_heading_styles(docx_path: Path, toc_entries: list["TocEntry"]) -> None:
    """Apply semantic heading styles to body paragraphs that match the PDF outline."""
    if not toc_entries:
        return

    from docx import Document

    document = Document(str(docx_path))
    title_to_level: dict[str, int] = {}
    for entry in toc_entries:
        key = normalize_lookup(entry.title)
        if not key:
            continue
        title_to_level[key] = min(title_to_level.get(key, 99), entry.level)

    for paragraph in document.paragraphs:
        text = _clean_text_fragment(paragraph.text)
        if not text:
            continue
        if "\t" in text or "w:hyperlink" in paragraph._p.xml:
            continue

        normalized = normalize_lookup(text)
        level = title_to_level.get(normalized)
        if level is None:
            continue

        cleaned_title = clean_toc_title(text)
        _rewrite_paragraph_text(paragraph, cleaned_title)
        try:
            paragraph.style = document.styles[f"Heading {min(level + 1, 4)}"]
        except KeyError:
            continue

    document.save(str(docx_path))


def apply_docx_page_geometry(
    docx_path: Path,
    page_width_pt: float | None = None,
    page_height_pt: float | None = None,
    page_sizes_pt: list[tuple[float, float]] | None = None,
) -> None:
    """Align DOCX sections to PDF page geometry, preserving portrait and landscape pages."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt

    document = Document(str(docx_path))

    sizes = page_sizes_pt or []
    if not sizes and page_width_pt is not None and page_height_pt is not None:
        sizes = [(page_width_pt, page_height_pt)]
    if not sizes:
        return

    for index, section in enumerate(document.sections):
        width_pt, height_pt = sizes[min(index, len(sizes) - 1)]
        section.orientation = WD_ORIENT.LANDSCAPE if width_pt > height_pt else WD_ORIENT.PORTRAIT
        section.page_width = Pt(width_pt)
        section.page_height = Pt(height_pt)

    document.save(str(docx_path))


def detect_pdf_page_profiles(
    page_sizes_pt: list[tuple[float, float]],
    *,
    tolerance_pt: float = 12.0,
) -> list[PdfPageProfile]:
    """
    Classify PDF page sizes into standard paper profiles when they are close
    enough to known sizes such as Letter, Legal, A4, or Tabloid.
    """
    profiles: list[PdfPageProfile] = []

    for width_pt, height_pt in page_sizes_pt:
        width_pt = float(width_pt)
        height_pt = float(height_pt)
        orientation = "landscape" if width_pt > height_pt else "portrait"
        short_side, long_side = sorted((width_pt, height_pt))

        best_name = "custom"
        best_size = (short_side, long_side)
        matched_standard = False
        best_delta = float("inf")

        for paper_name, (paper_short, paper_long) in _STANDARD_PAPER_SIZES_PT.items():
            delta = max(abs(short_side - paper_short), abs(long_side - paper_long))
            if delta < best_delta:
                best_delta = delta
                best_name = paper_name
                best_size = (paper_short, paper_long)
                matched_standard = delta <= tolerance_pt

        if matched_standard:
            normalized_short, normalized_long = best_size
        else:
            best_name = "custom"
            normalized_short, normalized_long = short_side, long_side

        if orientation == "landscape":
            normalized_width_pt, normalized_height_pt = normalized_long, normalized_short
        else:
            normalized_width_pt, normalized_height_pt = normalized_short, normalized_long

        profiles.append(
            PdfPageProfile(
                width_pt=width_pt,
                height_pt=height_pt,
                normalized_width_pt=float(normalized_width_pt),
                normalized_height_pt=float(normalized_height_pt),
                paper_size=best_name,
                orientation=orientation,
                matched_standard=matched_standard,
            )
        )

    return profiles


def apply_docx_section_paper_profiles(docx_path: Path, page_profiles: list[PdfPageProfile]) -> None:
    """
    Normalize DOCX section paper sizes to the PDF's detected paper profile when
    each orientation maps cleanly to a single page size.
    """
    from docx import Document
    from docx.shared import Pt

    if not page_profiles:
        return

    profiles_by_orientation: dict[str, list[tuple[float, float]]] = {"portrait": [], "landscape": []}
    for profile in page_profiles:
        profiles_by_orientation[profile.orientation].append(
            (
                round(profile.normalized_width_pt, 2),
                round(profile.normalized_height_pt, 2),
            )
        )

    representative_sizes: dict[str, tuple[float, float]] = {}
    for orientation, sizes in profiles_by_orientation.items():
        unique_sizes = set(sizes)
        if len(unique_sizes) == 1:
            representative_sizes[orientation] = next(iter(unique_sizes))

    if not representative_sizes:
        return

    document = Document(str(docx_path))
    changed = False
    for section in document.sections:
        orientation = "landscape" if float(section.page_width.pt) > float(section.page_height.pt) else "portrait"
        target = representative_sizes.get(orientation)
        if target is None:
            continue
        target_width_pt, target_height_pt = target
        if abs(float(section.page_width.pt) - target_width_pt) > 0.1:
            section.page_width = Pt(target_width_pt)
            changed = True
        if abs(float(section.page_height.pt) - target_height_pt) > 0.1:
            section.page_height = Pt(target_height_pt)
            changed = True

    if changed:
        document.save(str(docx_path))


def apply_docx_page_margins_from_pdf(docx_path: Path, pdf_path: Path) -> None:
    """Approximate Word section margins from the real PDF content box on each page."""
    import fitz
    from docx import Document
    from docx.shared import Pt

    document = Document(str(docx_path))
    sections = list(document.sections)
    if not sections:
        return

    pdf = fitz.open(str(pdf_path))
    try:
        for index, section in enumerate(sections):
            page = pdf[min(index, len(pdf) - 1)]
            xs: list[float] = []
            ys: list[float] = []
            xe: list[float] = []
            ye: list[float] = []

            for block in page.get_text("dict")["blocks"]:
                x0, y0, x1, y1 = [float(v) for v in block["bbox"]]
                width = x1 - x0
                height = y1 - y0
                if block.get("type") == 1 and width >= page.rect.width * 0.9 and height <= 40:
                    continue
                if block.get("type") not in {0, 1}:
                    continue
                xs.append(x0)
                ys.append(y0)
                xe.append(x1)
                ye.append(y1)

            if not xs:
                continue

            left_margin = max(18.0, min(xs))
            top_margin = max(18.0, min(ys))
            right_margin = max(18.0, float(page.rect.width) - max(xe))
            bottom_margin = max(18.0, float(page.rect.height) - max(ye))

            section.left_margin = Pt(left_margin)
            section.top_margin = Pt(top_margin)
            section.right_margin = Pt(right_margin)
            section.bottom_margin = Pt(bottom_margin)
            section.header_distance = Pt(0)
            section.footer_distance = Pt(0)
    finally:
        pdf.close()

    document.save(str(docx_path))


def clean_docx_control_chars(docx_path: Path) -> None:
    """Remove zero-width/control characters without altering semantic content."""
    from docx import Document

    document = Document(str(docx_path))

    for paragraph in document.paragraphs:
        _rewrite_paragraph_text(paragraph, paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _rewrite_paragraph_text(paragraph, paragraph.text)

    document.save(str(docx_path))


def split_docx_merged_toc_paragraphs(docx_path: Path) -> None:
    """Split a single merged TOC paragraph into one paragraph per visible entry."""
    from docx import Document

    document = Document(str(docx_path))
    paragraphs = list(document.paragraphs)
    toc_title_index = next(
        (
            idx
            for idx, paragraph in enumerate(paragraphs)
            if _clean_text_fragment(paragraph.text).lower() in {"indice", "índice"}
        ),
        None,
    )
    if toc_title_index is None or toc_title_index + 1 >= len(paragraphs):
        return

    toc_paragraph = paragraphs[toc_title_index + 1]
    entries = _extract_toc_lines(toc_paragraph.text)
    if len(entries) <= 1:
        return

    _rewrite_paragraph_text(toc_paragraph, entries[0])
    cursor = toc_paragraph
    for entry in entries[1:]:
        cursor = _insert_paragraph_after(cursor)
        _rewrite_paragraph_text(cursor, entry)

    document.save(str(docx_path))


def link_existing_docx_toc_lines(docx_path: Path) -> None:
    """Make existing TOC lines clickable while preserving their visible text."""
    from docx import Document

    document = Document(str(docx_path))
    paragraphs = list(document.paragraphs)
    if not paragraphs:
        return

    page_starts = _page_start_indices(paragraphs)
    page_bookmarks: dict[int, str] = {}
    bookmark_id = 1
    for page_number, paragraph_index in enumerate(page_starts, start=1):
        bookmark_name = f"folio_page_{page_number}"
        _append_bookmark(paragraphs[paragraph_index], bookmark_name, bookmark_id)
        page_bookmarks[page_number] = bookmark_name
        bookmark_id += 1

    toc_index = next(
        (
            idx
            for idx, paragraph in enumerate(paragraphs)
            if _clean_text_fragment(paragraph.text).lower() in {"indice", "índice"}
        ),
        None,
    )
    if toc_index is None:
        document.save(str(docx_path))
        return

    toc_pattern = re.compile(r"^(.*?\.{3,}\s*)(\d+)\s*$")
    toc_started = False
    for paragraph in paragraphs[toc_index + 1:]:
        text = _clean_text_fragment(paragraph.text)
        if not text:
            if toc_started:
                break
            continue
        match = toc_pattern.match(text)
        if not match:
            if toc_started:
                break
            continue
        toc_started = True
        page_number = int(match.group(2))
        target = page_bookmarks.get(page_number)
        if not target:
            continue
        _clear_paragraph_content(paragraph)
        _append_internal_hyperlink_plain(paragraph, text, target)

    document.save(str(docx_path))


def overlay_pdf_visual_layers(docx_path: Path, pdf_path: Path) -> None:
    """Overlay a visual non-text layer per PDF page while keeping DOCX text editable."""
    from docx import Document
    import fitz

    document = Document(str(docx_path))
    sections = list(document.sections)
    if not sections:
        return

    pdf = fitz.open(str(pdf_path))
    shape_id = 1000
    inserted = 0
    try:
        for page_index, page in enumerate(pdf):
            if page_index >= len(sections):
                continue
            image_bytes = _render_page_visual_layer(page)
            if not image_bytes:
                continue
            section = sections[page_index]
            header = section.header
            header.is_linked_to_previous = False
            header_paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            _clear_paragraph_content(header_paragraph)
            for extra_paragraph in list(header.paragraphs[1:]):
                _delete_paragraph(extra_paragraph)
            _append_floating_picture(
                header_paragraph,
                image_bytes=image_bytes,
                x_pt=0.0,
                y_pt=0.0,
                width_pt=float(page.rect.width),
                height_pt=float(page.rect.height),
                shape_id=shape_id,
                behind_text=True,
            )
            shape_id += 1
            inserted += 1
    finally:
        pdf.close()

    if inserted:
        document.save(str(docx_path))


def overlay_pdf_header_footer_bands(docx_path: Path, pdf_path: Path) -> None:
    """Overlay only repeated full-width decorative bands from the PDF into section headers."""
    from docx import Document
    import fitz

    def orientation_key(width: float, height: float) -> str:
        return "landscape" if width > height else "portrait"

    pdf = fitz.open(str(pdf_path))
    bands_by_orientation: dict[str, list[tuple[bytes, float, float, float, float]]] = {}
    try:
        for page in pdf:
            key = orientation_key(float(page.rect.width), float(page.rect.height))
            if key in bands_by_orientation:
                continue
            bands: list[tuple[bytes, float, float, float, float]] = []
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 1 or not block.get("image"):
                    continue
                x0, y0, x1, y1 = [float(v) for v in block["bbox"]]
                width = x1 - x0
                height = y1 - y0
                if width >= page.rect.width * 0.9 and height <= 40:
                    bands.append((block["image"], x0, y0, width, height))
            if bands:
                bands_by_orientation[key] = bands
    finally:
        pdf.close()

    if not bands_by_orientation:
        return

    document = Document(str(docx_path))
    shape_id = 6000
    inserted = 0
    for section in document.sections:
        key = orientation_key(float(section.page_width.pt), float(section.page_height.pt))
        bands = bands_by_orientation.get(key)
        if not bands:
            continue
        header = section.header
        header.is_linked_to_previous = False
        anchor_paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        for image_bytes, x_pt, y_pt, width_pt, height_pt in bands:
            _append_floating_picture(
                anchor_paragraph,
                image_bytes=image_bytes,
                x_pt=x_pt,
                y_pt=y_pt,
                width_pt=width_pt,
                height_pt=height_pt,
                shape_id=shape_id,
                behind_text=True,
            )
            shape_id += 1
            inserted += 1

    if inserted:
        document.save(str(docx_path))


def compact_trailing_body_whitespace(docx_path: Path, *, tail_window: int = 16) -> None:
    """
    Remove top-level whitespace-only paragraphs near the document tail that
    commonly appear after native Word PDF imports and can push trailing tables
    or a final blank page.

    The cleanup is intentionally conservative:
    - it only inspects the last `tail_window` body elements;
    - it never removes paragraphs that contain drawings;
    - it targets whitespace-only separators adjacent to tables or the final
      section properties.
    """
    from docx import Document

    document = Document(str(docx_path))
    body = document._element.body

    def _is_blank_body_paragraph(element) -> bool:
        if not element.tag.endswith("}p"):
            return False
        texts = "".join(element.xpath(".//w:t/text()"))
        drawings = element.xpath(".//w:drawing|.//w:pict")
        return not texts.strip() and not drawings

    changed = False
    while True:
        body_children = list(body)
        start_index = max(0, len(body_children) - tail_window)
        removed_in_pass = False

        for element in body_children[start_index:-1]:
            if not _is_blank_body_paragraph(element):
                continue

            previous = element.getprevious()
            next_element = element.getnext()
            previous_tag = previous.tag.split("}")[-1] if previous is not None else None
            next_tag = next_element.tag.split("}")[-1] if next_element is not None else None

            if previous_tag == "tbl" or next_tag in {"tbl", "sectPr"}:
                body.remove(element)
                changed = True
                removed_in_pass = True
                break

        if not removed_in_pass:
            break

    if changed:
        document.save(str(docx_path))


def compact_comment_guidelines(docx_path: Path) -> None:
    """Tighten the visual comment-guide block that Word often imports with extra whitespace."""
    from docx import Document
    from docx.shared import Pt

    document = Document(str(docx_path))

    def _is_blank_paragraph_element(element) -> bool:
        if not element.tag.endswith("}p"):
            return False
        texts = "".join(element.xpath(".//w:t/text()"))
        drawings = element.xpath(".//w:drawing|.//w:pict")
        return not texts.strip() and not drawings

    changed = False
    for paragraph in list(document.paragraphs):
        text = (paragraph.text or "").strip()
        if "7.3 Comentarios Generales" in text:
            fmt = paragraph.paragraph_format
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing = 1.0
            changed = True

            next_element = paragraph._p.getnext()
            while next_element is not None and _is_blank_paragraph_element(next_element):
                to_remove = next_element
                next_element = next_element.getnext()
                to_remove.getparent().remove(to_remove)
                changed = True
            continue

        if text and set(text) == {"_"}:
            fmt = paragraph.paragraph_format
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing = 1.0
            changed = True

            previous = paragraph._p.getprevious()
            if previous is not None and _is_blank_paragraph_element(previous):
                previous.getparent().remove(previous)
                changed = True

    if changed:
        document.save(str(docx_path))


def build_positioned_docx_from_pdf(docx_path: Path, pdf_path: Path, page_sizes_pt: list[tuple[float, float]]) -> None:
    """Build a DOCX using exact page visuals plus positioned editable text blocks."""
    import fitz
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.shared import Pt

    document = Document()
    first_anchor = document.add_paragraph()

    def configure_section(section, width_pt: float, height_pt: float) -> None:
        section.orientation = WD_ORIENT.LANDSCAPE if width_pt > height_pt else WD_ORIENT.PORTRAIT
        section.page_width = Pt(width_pt)
        section.page_height = Pt(height_pt)
        section.left_margin = Pt(0)
        section.right_margin = Pt(0)
        section.top_margin = Pt(0)
        section.bottom_margin = Pt(0)
        section.header_distance = Pt(0)
        section.footer_distance = Pt(0)

    configure_section(document.sections[0], *page_sizes_pt[0])
    anchor_paragraphs = [first_anchor]
    _clear_paragraph_content(first_anchor)

    for width_pt, height_pt in page_sizes_pt[1:]:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section, width_pt, height_pt)
        paragraph = document.add_paragraph()
        _clear_paragraph_content(paragraph)
        anchor_paragraphs.append(paragraph)

    temp_path = docx_path.with_name(f"{docx_path.stem}._layout_tmp.docx")
    document.save(str(temp_path))

    overlay_pdf_visual_layers(temp_path, pdf_path)

    document = Document(str(temp_path))
    pdf = fitz.open(str(pdf_path))
    shape_id = 4000
    try:
        for page_index, page in enumerate(pdf):
            if page_index >= len(document.paragraphs):
                break
            anchor = document.paragraphs[page_index]
            for line in _page_line_items(page):
                x0, y0, x1, y1 = line["bbox"]
                _append_positioned_textbox(
                    anchor,
                    x_pt=x0,
                    y_pt=y0,
                    width_pt=x1 - x0,
                    height_pt=y1 - y0,
                    line=line,
                    shape_id=shape_id,
                )
                shape_id += 1
    finally:
        pdf.close()

    document.save(str(docx_path))
    try:
        temp_path.unlink(missing_ok=True)
    except Exception:
        pass


def validate_docx_content(docx_path: Path) -> None:
    """Basic integrity check to catch empty or corrupted DOCX outputs early."""
    from docx import Document

    document = Document(str(docx_path))
    paragraph_count = sum(1 for p in document.paragraphs if p.text.strip())
    table_count = len(document.tables)
    inline_shapes = len(document.inline_shapes)
    textbox_shapes = sum(1 for p in document.paragraphs if "w:txbxContent" in p._p.xml or "w:pict" in p._p.xml)

    if paragraph_count == 0 and table_count == 0 and inline_shapes == 0 and textbox_shapes == 0:
        raise ValueError("DOCX generado sin contenido util")


def enhance_basic_styles(docx_path: Path) -> None:
    """Normalize control characters and convert extracted bullets into real Word lists."""
    from docx import Document

    document = Document(str(docx_path))

    for paragraph in list(document.paragraphs):
        original_style = (paragraph.style.name if paragraph.style else "") or ""
        if original_style.startswith("Heading"):
            _rewrite_paragraph_text(paragraph, paragraph.text)
            continue

        segments = _paragraph_segments(paragraph.text)
        if not segments:
            _rewrite_paragraph_text(paragraph, paragraph.text)
            continue

        prefix_segments: list[str] = []
        list_blocks: list[tuple[str, str | None, str]] = []
        for segment in segments:
            kind, marker, text = _parse_list_segment(segment)
            if kind == "text":
                if list_blocks:
                    previous_kind, previous_marker, previous_text = list_blocks[-1]
                    list_blocks[-1] = (
                        previous_kind,
                        previous_marker,
                        _clean_text_fragment(f"{previous_text} {text}"),
                    )
                else:
                    prefix_segments.append(text)
                continue
            list_blocks.append((kind, marker, text))

        if not list_blocks:
            _rewrite_paragraph_text(paragraph, "\n".join(segments))
            continue

        blocks: list[tuple[str, str | None, str]] = []
        if prefix_segments:
            blocks.append(("text", None, " ".join(prefix_segments)))
        blocks.extend(list_blocks)

        _apply_block(paragraph, document, *blocks[0])
        cursor = paragraph
        for block in blocks[1:]:
            cursor = _insert_paragraph_after(cursor)
            _apply_block(cursor, document, *block)

    document.save(str(docx_path))


def apply_native_list_styles_from_pdf(docx_path: Path, pdf_path: Path) -> None:
    """Convert literal bullets/numbering into native Word list styles using the PDF as source of truth."""
    from docx import Document

    document = Document(str(docx_path))
    pdf_candidates = _extract_pdf_native_list_candidates(pdf_path)

    for paragraph in document.paragraphs:
        text = _clean_text_fragment(paragraph.text)
        if not text or _looks_like_toc_entry(text):
            continue
        if (paragraph.style.name if paragraph.style else "").startswith("Heading"):
            continue

        kind, marker, stripped_text = _parse_list_segment(text)
        if kind != "text" and stripped_text:
            level = 1
            if kind == "bullet":
                if marker in {"○", "◦"}:
                    level = 2
                elif marker in {"▪", "■"}:
                    level = 3
            elif marker and marker.isalpha():
                level = 2
            _apply_native_list_style(paragraph, document, kind, level, stripped_text, rewrite_text=True)
            continue

        normalized = normalize_lookup(text)
        pdf_match = pdf_candidates.get(normalized)
        if not pdf_match:
            continue
        kind, level = pdf_match
        _apply_native_list_style(paragraph, document, kind, level, text, rewrite_text=False)

    document.save(str(docx_path))


def split_docx_toc_lines_from_pdf(docx_path: Path, pdf_path: Path) -> None:
    """Split a merged TOC paragraph using the PDF's own TOC lines as source of truth."""
    import fitz
    from docx import Document

    pdf = fitz.open(str(pdf_path))
    toc_lines: list[str] = []
    try:
        for page in pdf:
            found_title = False
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    text = _clean_text_fragment("".join(span.get("text", "") for span in line.get("spans", [])))
                    if not text:
                        continue
                    if text.lower() in {"indice", "índice"}:
                        found_title = True
                        continue
                    if found_title and _looks_like_toc_entry(text):
                        toc_lines.append(text)
                    elif found_title and toc_lines:
                        break
                if found_title and toc_lines and not _looks_like_toc_entry(toc_lines[-1]):
                    break
            if toc_lines:
                break
    finally:
        pdf.close()

    if len(toc_lines) <= 1:
        return

    document = Document(str(docx_path))
    paragraphs = list(document.paragraphs)
    toc_title_index = next(
        (
            idx
            for idx, paragraph in enumerate(paragraphs)
            if _clean_text_fragment(paragraph.text).lower() in {"indice", "índice"}
        ),
        None,
    )
    if toc_title_index is None or toc_title_index + 1 >= len(paragraphs):
        return

    target = paragraphs[toc_title_index + 1]
    if len(_extract_toc_lines(target.text)) <= 1 and len(_meaningful_lines(target.text)) <= 1:
        return

    _rewrite_paragraph_text(target, toc_lines[0])
    cursor = target
    for line in toc_lines[1:]:
        cursor = _insert_paragraph_after(cursor)
        _rewrite_paragraph_text(cursor, line)

    document.save(str(docx_path))
